from fastapi import FastAPI, HTTPException, Body
from pydantic import BaseModel
from typing import List
import httpx
import base64
from docx import Document
from datetime import datetime
import os
from fastapi.responses import StreamingResponse
import io
import cv2      # Importujemy OpenCV do przetwarzania obrazów
import numpy as np # Importujemy NumPy do operacji na tablicach

app = FastAPI(title="API Generatora Raportów")

# Poprawny adres URL serwisu archiwum, zgodny z docker-compose.yml
ARCHIWUM_API_URL = os.getenv("ARCHIWUM_API_URL", "http://archiwum:8000")

class BatchRequest(BaseModel):
    sample_ids: List[str]

@app.post("/generate-report")
async def generate_report(request: BatchRequest):
    print(f"--- SERWIS RAPORT: Otrzymano żądanie dla {len(request.sample_ids)} próbek ---")
    
    samples_data = []
    # Pobieramy dane dla każdej próbki osobno
    async with httpx.AsyncClient() as client:
        for sample_id in request.sample_ids:
            try:
                url = f"{ARCHIWUM_API_URL}/samples/{sample_id}"
                print(f"RAPORT: Pobieranie danych dla próbki '{sample_id}' z adresu: {url}")
                response = await client.get(url)
                response.raise_for_status()
                samples_data.append(response.json())
            except httpx.RequestError as e:
                print(f"!!! OSTRZEŻENIE: Nie udało się pobrać danych dla próbki '{sample_id}': {e}")
                continue

    if not samples_data:
        raise HTTPException(status_code=404, detail="Nie udało się pobrać danych dla żadnej z wybranych próbek.")

    try:
        # Tworzenie dokumentu na podstawie szablonu
        doc = Document("szablon.docx")
        doc.add_heading("Raport z analizy próbek", level=1)

        for sample in samples_data:
            doc.add_heading(f"Próbka {sample.get('sample_number', 'Brak numeru')}", level=2)
            doc.add_paragraph(f"Numer próbki: {sample.get('sample_number', 'Brak')}")
            doc.add_paragraph(f"Data dodania: {sample.get('date_added', 'Brak')}")
            doc.add_paragraph(f"Data predykcji: {sample.get('date_predicted', 'Brak')}")
            doc.add_paragraph(f"Algorytm: {sample.get('algorithm', 'Brak')}")
            doc.add_paragraph(f"Status: {sample.get('status', 'Brak')}")
            doc.add_paragraph(f"Pewność (confidence): {sample.get('confidence', 'Brak')}")

            # --- OSTATECZNA POPRAWKA LOGIKI OBRAZU ---
            image_data_b64 = sample.get("cropped_image")
            if image_data_b64:
                try:
                    # Krok 1: Dekodujemy dane base64 do surowych bajtów
                    decoded_bytes = base64.b64decode(image_data_b64)
                    
                    # Krok 2: Konwertujemy te bajty na tablicę NumPy, a potem z powrotem na obraz OpenCV
                    nparr = np.frombuffer(decoded_bytes, np.uint8)
                    img_np = cv2.imdecode(nparr, cv2.IMREAD_COLOR)

                    # Sprawdzamy, czy obraz został poprawnie zdekodowany
                    if img_np is None:
                        raise ValueError("cv2.imdecode zwróciło None. Niepoprawne dane obrazu.")

                    # Krok 3: Kodujemy obraz na nowo do POPRAWNEGO formatu PNG w pamięci
                    is_success, buffer = cv2.imencode(".png", img_np)
                    if not is_success:
                        raise ValueError("Nie udało się przekonwertować obrazu do formatu PNG")

                    # Krok 4: Tworzymy strumień w pamięci i dodajemy do dokumentu
                    image_stream = io.BytesIO(buffer)
                    doc.add_picture(image_stream)

                except Exception as img_e:
                    # Jeśli cokolwiek pójdzie nie tak, logujemy błąd i wstawiamy tekst zastępczy
                    print(f"RAPORT: Błąd dodawania obrazu dla próbki {sample.get('sample_number')}: {img_e}")
                    doc.add_paragraph("(Błąd podczas renderowania obrazu w raporcie)")
        
        # Zapisanie gotowego dokumentu do strumienia w pamięci
        memory_stream = io.BytesIO()
        doc.save(memory_stream)
        memory_stream.seek(0)

        # Przygotowanie i wysłanie pliku do użytkownika
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        filename = f"raport_{timestamp}.docx"
        headers = {'Content-Disposition': f'attachment; filename="{filename}"'}
        
        print(f"RAPORT: Pomyślnie wygenerowano plik '{filename}'. Wysyłanie do przeglądarki.")
        
        return StreamingResponse(
            memory_stream,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=headers
        )

    except FileNotFoundError:
        print("!!! KRYTYCZNY BŁĄD: Brak pliku 'szablon.docx' w kontenerze!")
        raise HTTPException(status_code=500, detail="Brak pliku szablonu na serwerze.")
    except Exception as e:
        print(f"!!! BŁĄD KRYTYCZNY: Błąd podczas generowania pliku docx: {e}")
        raise HTTPException(status_code=500, detail=f"Błąd generowania raportu: {e}")